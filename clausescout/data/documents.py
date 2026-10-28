"""Reading uploaded contracts (PDF, DOCX, TXT) into text plus page or paragraph segments."""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import docx
import pymupdf

from clausescout.grounding.location import Segment, paragraph_segments

SUPPORTED = (".pdf", ".docx", ".txt")


@dataclass
class Document:
    name: str
    kind: str  # "pdf", "docx" or "txt"
    text: str
    segments: list[Segment]

    @property
    def looks_scanned(self) -> bool:
        # A PDF with pages but almost no text is probably a scan, which I can't read without OCR.
        return self.kind == "pdf" and len(self.text.strip()) < 50 * max(len(self.segments), 1)


def _join(blocks: list[str], kind: str) -> tuple[str, list[Segment]]:
    # I join blocks with a blank line and remember where each one sits in the full text.
    parts, segments, pos = [], [], 0
    for block in blocks:
        if kind == "paragraph" and not block.strip():
            continue
        segments.append(Segment(kind, len(segments) + 1, pos, pos + len(block)))
        parts.append(block)
        pos += len(block) + 2
    return "\n\n".join(parts), segments


def _pdf(data: bytes) -> tuple[str, list[Segment]]:
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        # Pages keep their real numbers even when one is blank, so citations stay right.
        pages = [page.get_text("text").strip() for page in pdf]
    return _join(pages, "page")


def _docx(data: bytes) -> tuple[str, list[Segment]]:
    d = docx.Document(io.BytesIO(data))
    blocks = [p.text for p in d.paragraphs]
    # Tables often hold key terms (fees, caps, notice periods), so I read each row as a block.
    for table in d.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                # Merged cells repeat, so I skip a cell identical to the one before it.
                if not cells or cell.text != cells[-1]:
                    cells.append(cell.text)
            blocks.append(" | ".join(cells))
    return _join(blocks, "paragraph")


def _txt(data: bytes) -> tuple[str, list[Segment]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    return text, paragraph_segments(text)


def load_document(source: str | Path | bytes, filename: str | None = None) -> Document:
    """Loads a contract from a path, or from raw bytes plus the original filename."""
    if isinstance(source, bytes):
        if filename is None:
            raise ValueError("A filename is needed to tell the file type of raw bytes")
        data, name = source, filename
    else:
        path = Path(source)
        data, name = path.read_bytes(), path.name
    suffix = Path(name).suffix.lower()
    readers = {".pdf": _pdf, ".docx": _docx, ".txt": _txt}
    if suffix not in readers:
        raise ValueError(f"Unsupported file type {suffix!r}; use one of {', '.join(SUPPORTED)}")
    text, segments = readers[suffix](data)
    return Document(name, suffix[1:], text, segments)
