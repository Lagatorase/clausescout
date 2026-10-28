"""Exporting reviewed contracts as a due diligence report: an Excel workbook and an HTML summary."""

from __future__ import annotations

import html
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from clausescout.data.cuad import CATEGORIES, HIGH_STAKES
from clausescout.findings import Finding
from clausescout.grounding.location import Segment, locate
from clausescout.playbook.engine import Flag

DISCLAIMER = (
    "ClauseScout is a review aid, not legal advice. It can miss clauses and mislabel them. "
    "Every finding quotes the contract word for word so it can be checked against the source."
)
SEVERITY_FILL = {"high": "F8D7DA", "medium": "FFF3CD", "low": "E2E3E5"}


@dataclass
class ContractReview:
    name: str
    findings: list[Finding]  # already thresholded and grounded
    flags: list[Flag] = field(default_factory=list)
    segments: list[Segment] = field(default_factory=list)

    def location(self, f: Finding) -> str:
        if f.start is None or f.end is None:
            return ""
        loc = locate(self.segments, f.start, f.end)
        return loc.label if loc else ""

    def missing(self) -> list[str]:
        found = {f.category for f in self.findings}
        return [c for c in CATEGORIES if c not in found]

    def sorted_findings(self) -> list[Finding]:
        # I keep CUAD's category order, then document order inside a category.
        order = {c: i for i, c in enumerate(CATEGORIES)}
        return sorted(self.findings, key=lambda f: (order.get(f.category, 99), f.start or 0))


def _evidence(review: ContractReview, flag: Flag) -> str:
    return "\n".join(f'"{e.text}" ({review.location(e)})' for e in flag.evidence)


def write_excel(reviews: list[ContractReview], path: str | Path) -> Path:
    """One workbook for a whole pack: summary, findings, flags and clause types not found."""
    wb = Workbook()
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")

    def sheet(title, header, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
        return ws

    ws = sheet("Summary", ["Contract", "High flags", "Medium flags", "Low flags", "Findings",
                           "High-stakes types found"], [50, 12, 14, 12, 10, 24], first=True)
    for r in reviews:
        sev = [fl.rule.severity for fl in r.flags]
        hs = len({f.category for f in r.findings} & set(HIGH_STAKES))
        ws.append([r.name, sev.count("high"), sev.count("medium"), sev.count("low"),
                   len(r.findings), f"{hs} of {len(HIGH_STAKES)}"])
    ws.append([])
    ws.append([f"Generated {date.today().isoformat()}. {DISCLAIMER}"])

    ws = sheet("Flags", ["Contract", "Severity", "Flag", "Explanation", "Evidence"],
               [40, 10, 36, 60, 80])
    for r in reviews:
        for fl in r.flags:
            ws.append([r.name, fl.rule.severity, fl.rule.title, fl.rule.explanation,
                       _evidence(r, fl)])
            fill = SEVERITY_FILL[fl.rule.severity]
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fill)

    ws = sheet("Findings", ["Contract", "Clause type", "High stakes", "Quote", "Location",
                            "Confidence", "Method"], [40, 30, 11, 90, 12, 11, 10])
    for r in reviews:
        for f in r.sorted_findings():
            ws.append([r.name, f.category, "yes" if f.category in HIGH_STAKES else "",
                       f.text, r.location(f), round(f.score, 3), f.method])

    ws = sheet("Not found", ["Contract", "Clause type", "High stakes"], [40, 34, 11])
    for r in reviews:
        for c in r.missing():
            ws.append([r.name, c, "yes" if c in HIGH_STAKES else ""])

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


_CSS = """
:root { --bg: #fff; --fg: #1d1f23; --muted: #5c6370; --line: #d8dbe0; --quote: #f4f5f7;
        --high: #b42318; --medium: #a15c07; --low: #4b5563; }
@media (prefers-color-scheme: dark) {
  :root { --bg: #16181c; --fg: #e6e7ea; --muted: #9aa0aa; --line: #2e3238; --quote: #202329;
          --high: #f97066; --medium: #fdb022; --low: #9aa0aa; }
}
body { background: var(--bg); color: var(--fg); font: 15px/1.5 system-ui, sans-serif;
       max-width: 960px; margin: 0 auto; padding: 24px 16px; }
h1 { font-size: 1.5rem; margin-bottom: 4px; } h2 { margin-top: 2.2rem; } h3 { margin-bottom: 4px; }
.muted { color: var(--muted); } .note { border-left: 3px solid var(--line); padding-left: 12px; }
table { border-collapse: collapse; width: 100%; } td, th { text-align: left; padding: 6px 8px;
        border-bottom: 1px solid var(--line); vertical-align: top; }
.sev { font-weight: 600; text-transform: capitalize; }
.high { color: var(--high); } .medium { color: var(--medium); } .low { color: var(--low); }
blockquote { background: var(--quote); margin: 6px 0; padding: 8px 12px; border-radius: 4px; }
"""


def render_html(reviews: list[ContractReview]) -> str:
    """A self-contained HTML summary, readable on screen and printable to PDF."""
    e = html.escape
    parts = [f"<h1>Due diligence summary</h1><p class='muted'>{len(reviews)} contract(s), "
             f"generated {date.today().isoformat()}</p><p class='note'>{e(DISCLAIMER)}</p>"]
    rows = []
    for r in reviews:
        sev = [fl.rule.severity for fl in r.flags]
        rows.append(f"<tr><td>{e(r.name)}</td>" + "".join(
            f"<td class='{s}'>{sev.count(s)}</td>" for s in ("high", "medium", "low")) +
            f"<td>{len(r.findings)}</td></tr>")
    parts.append("<table><tr><th>Contract</th><th>High</th><th>Medium</th><th>Low</th>"
                 f"<th>Findings</th></tr>{''.join(rows)}</table>")
    for r in reviews:
        parts.append(f"<h2>{e(r.name)}</h2><h3>Flags</h3>")
        if not r.flags:
            parts.append("<p class='muted'>No playbook flags.</p>")
        for fl in r.flags:
            quotes = "".join(f"<blockquote>{e(ev.text)} <span class='muted'>"
                             f"({e(r.location(ev))})</span></blockquote>" for ev in fl.evidence)
            parts.append(f"<p><span class='sev {fl.rule.severity}'>{fl.rule.severity}</span> "
                         f"<b>{e(fl.rule.title)}</b><br>{e(fl.rule.explanation)}</p>{quotes}")
        found = {f.category for f in r.findings}
        hs_rows = []
        for c in HIGH_STAKES:
            quotes = [f for f in r.sorted_findings() if f.category == c]
            cell = "".join(f"<blockquote>{e(f.text)} <span class='muted'>({e(r.location(f))}, "
                           f"{f.score:.2f})</span></blockquote>" for f in quotes)
            hs_rows.append(f"<tr><td>{e(c)}</td><td>{cell or '<span class=muted>Not found</span>'}"
                           "</td></tr>")
        parts.append("<h3>High-stakes clauses</h3><table>" + "".join(hs_rows) + "</table>")
        other = sorted(found - set(HIGH_STAKES), key=CATEGORIES.index)
        parts.append(f"<p class='muted'>Other clause types found: {e(', '.join(other)) or 'none'}."
                     " The Excel workbook lists every quote.</p>")
    return ("<!doctype html><html lang='en'><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width, initial-scale=1'>"
            f"<title>Due diligence summary</title><style>{_CSS}</style></head><body>"
            + "".join(parts) + "</body></html>")


def write_html(reviews: list[ContractReview], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_html(reviews), encoding="utf-8")
    return path
