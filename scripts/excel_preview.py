"""Render the first rows of an Excel report's sheets as an HTML page, for the README screenshot.

The app only offers the Excel file as a download, so for the README I show what it contains:
the header and first rows of the chosen sheets, as plain tables. web/scripts/shoot-page.mjs
then takes the screenshot.

Usage: python scripts/excel_preview.py <report.xlsx> <preview.html> [--sheets Flags Findings]
"""

import argparse
import html
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS, MAX_CHARS = 8, 110


def table(ws) -> str:
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
        if len(rows) > MAX_ROWS:
            break

    def cell(v: str, tag: str) -> str:
        # Long quotes are shortened so each table fits on one screen.
        short = v if len(v) <= MAX_CHARS else v[:MAX_CHARS] + "..."
        return f"<{tag}>{html.escape(short)}</{tag}>"

    head = "".join(cell(v, "th") for v in rows[0]) if rows else ""
    body = "".join("<tr>" + "".join(cell(v, "td") for v in r) + "</tr>" for r in rows[1:])
    return (f"<h2>Sheet: {html.escape(ws.title)} (first {len(rows) - 1} rows)</h2>"
            f"<table><tr>{head}</tr>{body}</table>")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--sheets", nargs="+", default=["Flags", "Findings"])
    args = ap.parse_args()

    wb = load_workbook(args.src, read_only=True)
    names = [s.title for s in wb.worksheets]
    tables = "".join(table(wb[n]) for n in args.sheets if n in names)
    args.out.write_text(f"""<!doctype html><html><head><meta charset="utf-8"><style>
body {{ font: 13px/1.4 Segoe UI, sans-serif; margin: 20px; color: #1d1d1f; background: #fff; }}
h1 {{ font-size: 16px; margin: 0 0 4px; }} h2 {{ font-size: 14px; margin: 18px 0 6px; }}
p {{ color: #666; margin: 0 0 6px; }} table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 1px solid #d0d0d0; padding: 4px 6px; text-align: left; vertical-align: top; }}
th {{ background: #1f4e78; color: #fff; }} tr:nth-child(even) td {{ background: #f3f6fa; }}
</style></head><body><h1>Excel report preview</h1>
<p>Sheets in the file: {html.escape(", ".join(names))}.</p>{tables}</body></html>""",
                        encoding="utf-8")
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
