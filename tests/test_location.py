"""Tests for document loading and mapping spans to pages or paragraphs."""

import docx
import pymupdf

from clausescout.data.documents import load_document
from clausescout.grounding.location import Segment, locate, paragraph_segments
from clausescout.grounding.verbatim import GroundingIndex


def test_paragraph_segments_number_non_empty_blocks():
    text = "Title\n\n\n1. Term. Two years.\n  \n2. Law. New York.\nstill clause two"
    segs = paragraph_segments(text)
    assert [s.number for s in segs] == [1, 2, 3]
    assert text[segs[2].start:segs[2].end] == "2. Law. New York.\nstill clause two"


def test_locate_single_and_spanning_segments():
    segs = [Segment("page", 1, 0, 10), Segment("page", 2, 12, 30), Segment("page", 3, 32, 50)]
    assert locate(segs, 3, 8).label == "p. 1"
    assert locate(segs, 15, 40).label == "pp. 2-3"
    # A span starting in the gap between pages belongs to the next page.
    assert locate(segs, 11, 20).label == "p. 2"
    para = [Segment("paragraph", 4, 0, 10), Segment("paragraph", 5, 11, 20)]
    assert locate(para, 2, 15).label == "paras. 4-5"
    assert locate(para, 12, 15).label == "para. 5"


def test_txt_document(tmp_path):
    p = tmp_path / "nda.txt"
    p.write_text("MUTUAL NDA\n\n1. Term.\n\n2. This Agreement is governed by Singapore law.",
                 encoding="utf-8")
    doc = load_document(p)
    assert doc.kind == "txt"
    start, end = GroundingIndex(doc.text).find("governed by Singapore law")
    assert locate(doc.segments, start, end).label == "para. 3"


def test_pdf_document_cites_pages(tmp_path):
    p = tmp_path / "msa.pdf"
    pdf = pymupdf.open()
    for body in ["Master Services Agreement\n1. Services.", "7. Governing Law.\nSingapore law."]:
        page = pdf.new_page()
        page.insert_text((72, 72), body)
    pdf.save(p)
    doc = load_document(p)
    assert doc.kind == "pdf" and [s.number for s in doc.segments] == [1, 2]
    start, end = GroundingIndex(doc.text).find("7. Governing Law. Singapore law.")
    assert locate(doc.segments, start, end).label == "p. 2"


def test_docx_document_cites_paragraphs_and_reads_tables(tmp_path):
    p = tmp_path / "supply.docx"
    d = docx.Document()
    d.add_paragraph("SUPPLY AGREEMENT")
    d.add_paragraph("")
    d.add_paragraph("1. Exclusivity. Supplier shall sell only to Buyer.")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Liability cap"
    table.cell(0, 1).text = "Fees paid in the prior 12 months"
    d.save(p)
    doc = load_document(p)
    assert doc.kind == "docx"
    start, end = GroundingIndex(doc.text).find("Supplier shall sell only to Buyer.")
    # The empty paragraph is skipped, so this is the second numbered paragraph.
    assert locate(doc.segments, start, end).label == "para. 2"
    assert GroundingIndex(doc.text).find("Fees paid in the prior 12 months") is not None


def test_load_from_bytes(tmp_path):
    doc = load_document(b"Clause one.\n\nClause two.", filename="x.txt")
    assert doc.name == "x.txt" and len(doc.segments) == 2
