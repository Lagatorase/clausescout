"""The Excel and HTML report export on a small made-up review."""

from openpyxl import load_workbook

from clausescout.findings import Finding
from clausescout.grounding.location import paragraph_segments
from clausescout.playbook.engine import apply_playbook, load_playbook
from clausescout.report.export import ContractReview, render_html, write_excel

TEXT = ("Supply Agreement\n\nThis Agreement is governed by the laws of New York.\n\n"
        "Supplier shall sell only to Buyer <b>.")


def _review():
    quote = "This Agreement is governed by the laws of New York."
    start = TEXT.index(quote)
    excl = "Supplier shall sell only to Buyer <b>."
    findings = [Finding("c", "Governing Law", quote, 0.8, start, start + len(quote), "qa"),
                Finding("c", "Exclusivity", excl, 0.7, TEXT.index(excl), len(TEXT), "qa")]
    flags = apply_playbook(load_playbook(), findings, TEXT)
    return ContractReview("supply.pdf", findings, flags, paragraph_segments(TEXT))


def test_excel_has_every_sheet_and_locations(tmp_path):
    path = write_excel([_review()], tmp_path / "r.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Flags", "Findings", "Not found"]
    rows = list(wb["Findings"].iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == "Governing Law" and rows[0][4] == "para. 2"
    assert len(list(wb["Not found"].iter_rows(min_row=2))) == 41 - 2
    flags = [r[2] for r in wb["Flags"].iter_rows(min_row=2, values_only=True)]
    assert "Exclusivity with no end date found" in flags


def test_html_escapes_quotes_and_says_it_is_not_legal_advice():
    page = render_html([_review()])
    assert "not legal advice" in page
    assert "&lt;b&gt;" in page and "<b>." not in page
    assert "para. 3" in page
